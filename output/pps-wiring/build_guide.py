from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(0.63)
sec.left_margin = sec.right_margin = Inches(0.70)
sec.header_distance = sec.footer_distance = Inches(0.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s=doc.styles[name]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=False
doc.core_properties.title='ZED F9P to USRP X300 PPS wiring'
doc.core_properties.subject='Physical assembly and bench qualification of a GPS PPS connection'
doc.core_properties.author='Prepared with Codex for the GNSS laboratory'
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('F9P to X300 PPS wiring  |  ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def p(text=''):
    for citation in ['[1, 2]', '[1, 5]', '[3, 5]']:
        text=text.replace(citation,citation.replace(' ','\u00a0'))
    return doc.add_paragraph(text)
def h(text): doc.add_heading(text,2)
def page(title):
    doc.add_page_break(); doc.add_heading(title,1)
def step(n,text):
    x=p(); x.add_run(str(n)+'. ').bold=True; x.add_run(text)
    x.paragraph_format.left_indent=Inches(.20)
    x.paragraph_format.first_line_indent=Inches(-.20)
def pic(filename,width,caption,alt):
    x=p(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_after=Pt(3)
    shape=x.add_run().add_picture(str(OUT/'images'/filename),width=Inches(width))
    shape._inline.docPr.set('descr',alt)
    doc.add_paragraph(caption,'Caption')
def table(headers,rows,widths,size=9.5):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for i,txt in enumerate(headers): t.rows[0].cells[i].text=txt
    for row in rows:
        cells=t.add_row().cells
        for i,txt in enumerate(row): cells[i].text=str(txt)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','5'),('color','D9D9D9')]: el.set(qn('w:'+k),v)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margins.append(el)
            tcp.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE7EF' if ri==0 else ('F4F6F8' if ri%2==0 else 'FFFFFF'));tcp.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.03
                if headers[i] in ['Quantity','Pin','Value']:
                    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for rr in para.runs: rr.font.size=Pt(size);rr.bold=(ri==0)
    x=p();x.paragraph_format.space_after=Pt(1);x.paragraph_format.space_before=Pt(0);x.paragraph_format.line_spacing=Pt(6)
def code(text):
    x=p(); x.paragraph_format.line_spacing=1.0
    r=x.add_run(text);r.font.name='Courier New';r.font.size=Pt(8.8)
def link(label,url):
    x=p();hl=OxmlElement('w:hyperlink');hl.set(qn('r:id'),x.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'0563C1');rp.append(co);r.append(rp)
    text=OxmlElement('w:t');text.text=label;r.append(text);hl.append(r);x._p.append(hl)

doc.add_paragraph('ZED F9P to USRP X300\nPPS wiring','Title')
p('Physical assembly guide for the photographed ZED-F9P-04B-01 carrier board. Prepared 23 September 2026.')
p('Connect the receiver PPS signal through a non-inverting 5 V buffer to the X300 rear SMA labelled PPS/TRIG IN. Connect their grounds through the coax shield. Keep any existing 10 MHz source on REF IN. [1, 2]')
pic('x300-pps-connectors.png',6.65,'Figure 1. GPT Image connector illustration. Positions are schematic, not the actual rear-panel layout. Identify the socket by its printed name.','Schematic X300 port guide: buffered PPS goes to PPS/TRIG IN; REF IN takes the separate 10 MHz clock. Coax centre carries PPS and shield carries ground.')
h('What this guide establishes')
p('The steps build and test a physical one-pulse-per-second input. The buffer circuit is a proposed bench design, not a measured or vendor-certified assembly. Scope checks are required on your actual board and X300 before relying on it.')
p('PPS gives the second boundary. It does not supply the continuous 10 MHz reference or tell the X300 the GPS week and time of week. Those remain separate connections and software work. [1, 5]')

page('1 Identify the F9P connection points')
pic('f9p-annotated.png',5.75,'Figure 2. GPT Image annotation of your supplied photo. Follow the visible PPS and GND silkscreen; do not count pads from a rotated board.','Actual F9P board annotation: PPS is on the left edge between RTK and RST; GND is the top left pad above 3V3; USB-C is at the top; SMA antenna connector is at the bottom.')
table(['Point','What to connect'],[
('PPS','A short signal wire to the buffer input. The PPS pad is between RTK and RST.'),
('GND','A short ground wire to the buffer ground. Use the left GND pad immediately above 3V3.'),
('USB-C','A data-capable USB cable for normal board power, configuration and GNSS messages.'),
('Bottom SMA','A compatible GNSS antenna. This antenna socket does not provide PPS.')],[1.05,5.72])
p('The module TIMEPULSE output is pin 53, but use the labelled carrier pad; do not solder to the module itself. Carrier routing, antenna bias and PPS voltage still need checking on this specific board. [2]')

page('2 Gather the parts and tools')
p('Use the buffered build below for the first installation. Its high-impedance input places very little load on the F9P. The output drives the coax from a separate 5 V supply. The exact carrier PPS drive circuit and X300 load have not been measured.')
table(['Item','Quantity','Selection'],[
('F9P and antenna','1 each','Your board, USB data cable, compatible GNSS antenna with a clear sky view.'),
('TC4427A driver','1','TC4427AEPA, 8-pin PDIP, or the same TC4427A function on a labelled adapter. Do not substitute the inverting TC4426A.'),
('Small solder board','1','Perfboard or prototyping PCB with a ground plane; 8-pin socket optional. Avoid loose breadboard wiring for the final assembly.'),
('Regulated supply','1','5.0 V for the buffer; a current-limited bench supply is convenient. Do not use the X300 12 V adapter.'),
('Decoupling capacitors','2','100 nF ceramic and 1 microfarad ceramic, at least 10 V rating, close to the driver supply pins.'),
('Series resistor','1','47 ohm, 0.25 W, between buffer OUT A and the coax centre.'),
('Input pull-down','1','47 kilohm between buffer IN A and GND.'),
('Short hookup wire','2','Flexible insulated 26-30 AWG; orange for PPS and black for GND. Keep the pair about 5-10 cm.'),
('Coax cable','1','Short flexible 50 ohm RG174 or similar; approximately 0.3-1 m is a practical starting length.'),
('X300 cable end','1','Standard SMA male with a centre pin; not RP-SMA. Prefer a factory-terminated cable or pigtail.'),
('Buffer cable end','1','SMA female bulkhead pigtail or PCB jack. Use a male-to-male cable to the X300.'),
('Assembly tools','Set','Fine-tip iron, electronics solder, flux, tweezers, magnifier, wire stripper, heat-shrink and strain relief.'),
('Test equipment','Set','Multimeter and oscilloscope with high-impedance x10 probes. A short SMA T can help measure at the X300 input.')],[1.25,.68,4.84],9.3)
h('Power arrangement')
p('Power the F9P through USB-C. Power the buffer from a separate regulated 5.0 V source. Join only their grounds and the PPS signal; do not join the external 5 V rail to the F9P 3V3 pad or to an unverified carrier power pad.')
p('The TC4427A input must not be driven while its supply is off. Turn the buffer on before the F9P, and turn the F9P off before the buffer. [4]')

page('3 Solder the F9P leads and prepare the coax')
step(1,'Stop the experiment. Unplug the F9P USB cable and buffer supply, and disconnect the X300 power and timing cables. Leave the soldering work on a clear bench.')
step(2,'Orient the F9P as in Figure 2: USB-C at the top and the large antenna SMA at the bottom. Locate PPS between RTK and RST, and the GND pad above 3V3. Check the labels again under magnification.')
step(3,'Cut the orange and black wires to about 5-10 cm. Strip only about 2 mm at each end, twist the strands and tin them lightly. Slide on any heat-shrink before soldering.')
step(4,'Apply a little flux and a small amount of solder to PPS and GND. Hold the tinned wire on the pad and heat only until the solder flows. Let it cool without movement. Do not repeatedly heat or pull on the pad.')
step(5,'Solder orange to PPS and black to GND. Inspect the gaps to RTK, RST and 3V3. Remove stray strands or solder bridges before power is applied.')
step(6,'Route the pair together toward the buffer. Secure the insulated wires to a standoff or enclosure, leaving a small service loop. Cable tension must not pull on the solder pads. Keep the backup battery and antenna components accessible.')
step(7,'Use a factory SMA pigtail at the buffer. Identify its centre conductor and braided shield with a multimeter. The centre conductor is the signal; the braid and SMA shell are ground.')
step(8,'At the buffer end, keep exposed coax conductors short, ideally only a few millimetres. Insulate the centre conductor from the braid. Fasten the connector or cable mechanically before attaching the electrical ends.')
h('Check continuity with all power disconnected')
table(['Measurement','Expected result'],[
('Orange wire to F9P PPS pad','Continuity.'),
('Black wire to F9P GND pad','Continuity.'),
('Bare cable centre pin to centre conductor','Continuity, before the cable is attached to the circuit.'),
('Bare cable shell to braid','Continuity.'),
('Bare cable centre to shell','Open circuit. Test the cable separately from electronics.'),
('Neighbouring board pads','No solder bridge. In-circuit semiconductor paths can affect a continuity reading; a beep alone is not proof of a short.')],[2.72,4.05])
p('Do not use resistance or continuity mode on a powered circuit. The SMA antenna socket, 3V3 pad, RTK pad and RST pad are not substitutes for PPS or GND.')

page('4 Assemble the buffer')
p('This proposed circuit uses one non-inverting channel of a TC4427A. It accepts a valid 3.3 V logic pulse at a high-impedance input and produces a pulse near its 5 V supply level. The IC is also sold as a MOSFET driver; here it is used as a line driver. [4]')
code('F9P PPS ----+---- IN A       OUT A ---- 47 ohm ---- SMA centre\n            |       TC4427A\n          47 kohm\n            |\nF9P GND ----+---- GND ----------------------------- SMA shield\n\nRegulated +5.0 V ---- VDD\n100 nF and 1 uF: each connected between VDD and GND')
p('For an 8-pin PDIP, look at the top with the notch upward. Pin 1 is upper left; count down the left side to 4, then up the right side from 5 to 8. The table is for PDIP/MSOP/SOIC, not the DFN package. [4]')
table(['Pin','Function','Connection'],[
('1','NC','Leave unconnected.'),
('2','IN A','Orange wire from F9P PPS; also 47 kilohm to pin 3.'),
('3','GND','F9P black wire, supply negative, coax shield, capacitor grounds and pin 4.'),
('4','IN B','Tie directly to pin 3; unused input must not float.'),
('5','OUT B','Leave unconnected.'),
('6','VDD','Regulated +5.0 V; both decoupling capacitors to pin 3.'),
('7','OUT A','47 ohm series resistor, then coax centre conductor.'),
('8','NC','Leave unconnected.')],[.45,.8,5.52])
step(9,'Fit the IC or socket on the solder board. Check orientation before soldering. Make the supply and ground paths short. Place the 100 nF capacitor immediately beside pins 6 and 3, with the 1 microfarad nearby.')
step(10,'Wire the table exactly. Put the 47 ohm resistor near pin 7, at the source end of the coax. Its purpose is to reduce reflections; do not add a 50 ohm resistor across the F9P PPS pad.')
step(11,'Mount the board inside a small insulated enclosure or on standoffs. Secure the SMA connector. Inspect solder joints and verify the +5 V rail has no hard short to GND.')
p('Do not treat the 1.5 A headline rating as a guaranteed continuous current at 5 V. Qualify this assembly using the measurements in Section 6. Its propagation delay is additional timing delay and must be measured for precision work.')

page('5 Configure one pulse per second')
p('Attach the GNSS antenna before applying USB power. Use a compatible antenna and the carrier manufacturer\'s antenna-bias arrangement; the photo does not establish the active-antenna supply. Place the antenna where it can receive the sky.')
p('Before powering the F9P, perform the buffer supply-only check in Step 12. Then power the buffer before the F9P. Connect a receiver configuration utility over USB, for example u-center on Windows or another UBX tool. Read the firmware version and save the existing configuration before changing timepulse settings.')
p('For HPG 1.32, the following CFG-TP values define a 1 Hz GPS-grid pulse with a 100 ms high interval. These are proposed setup values, not a read-back from your receiver. Match the interface documentation to your installed firmware. [3]')
table(['CFG-TP item','Value','Purpose'],[
('TP1_ENA','1','Enable timepulse.'),
('PULSE_DEF','0 / PERIOD','Use periods, in microseconds.'),
('PERIOD_TP1','1000000','One-second unlocked period.'),
('PERIOD_LOCK_TP1','1000000','One-second locked period.'),
('PULSE_LENGTH_DEF','1 / LENGTH','Use pulse length, in microseconds.'),
('LEN_TP1','0','No pulse in the unlocked parameter set.'),
('LEN_LOCK_TP1','100000','100 ms high pulse when using locked settings.'),
('SYNC_GNSS_TP1','1','Synchronize to GNSS time.'),
('USE_LOCKED_TP1','1','Use the locked settings when available.'),
('ALIGN_TO_TOW_TP1','1','Align pulse to the second boundary.'),
('POL_TP1','1','Rising edge marks the second.'),
('TIMEGRID_TP1','1 / GPS','Use the GPS time grid.')],[2.60,1.20,2.97],9.1)
p('Each first-column name has the prefix CFG-TP-. Read the values back, save them to supported nonvolatile storage, and check them again after a power cycle. Record any existing cable-delay or user-delay correction before changing it. Do not invent compensation values.')
h('Confirm timing validity')
p('Wait for valid receiver time and inspect its GNSS status. RTK corrections are not needed for basic PPS. A blinking PPS LED alone does not establish GPS synchronization. During signal loss, a receiver may retain estimated valid time for a while; zero unlocked pulse length is not an immediate antenna-loss detector. [3, 5]')
p('Keep USB/UART data available for time labels and validity checks. UBX-TIM-TP describes the next timepulse; use its documented time base and validity semantics when binding an epoch later. [5]')

page('6 Measure the signal before connecting the X300')
p('Keep the X300 cable disconnected for the first measurements. Use DC coupling and high-impedance x10 probes. Start near 200 ms per division to see the period, then zoom in on the rising edge. A scope set to 50 ohm input would heavily load the raw F9P output.')
step(12,'With the F9P disconnected, power the buffer at 5.0 V and verify pin 6 relative to pin 3 with a meter. Confirm the IC is not heating and the output is low with its input held down.')
step(13,'Connect the F9P ground and PPS leads with power removed, then power the buffer before the F9P. Measure the F9P PPS pad relative to its GND using a short probe ground connection.')
step(14,'Check the buffer output and then the far end of the intended coax. Use a suitable SMA measurement adapter. Do not short the centre pin with a large probe tip.')
table(['Check','Target for this build'],[
('Receiver PPS level','Low below 0.8 V; high comfortably above 2.4 V, normally near the carrier logic rail. Stop if high is marginal for the driver input.'),
('Rate and width','About 1.000 s between rising edges and about 100 ms high, matching Section 5.'),
('Buffered pulse','Same polarity and period. High near 5 V; low near 0 V. Confirm the pulse is not inverted.'),
('Edges','One clean transition; no substantial overshoot, undershoot or ringing that creates extra crossings.'),
('Receiver loading','Adding the powered buffer should not collapse or noticeably distort the F9P pulse.'),
('Output measurement','Record actual high/low levels, period, pulse width, cable length and probe setup.')],[1.55,5.22])
h('Load and voltage checks')
p('The F9P-04B data sheet specifies only 4 mA TIMEPULSE drive/sink. Do not connect a 50 ohm terminator directly to it: 3.3 V / 50 ohm would require about 66 mA. A 50 ohm coax cable is not the same as adding a 50 ohm DC termination. [2]')
p('Ettus specifies a 5 Vpp PPS waveform. This guide uses a unipolar pulse from 0 V to about +5 V, not a bipolar -2.5 V to +2.5 V signal. Use that as the output target. [1]')
p('The 47 ohm output resistor assumes the final input does not present a continuous 50 ohm DC load. A 50 ohm load would divide the pulse to roughly half its voltage. The public documentation inspected does not settle the load for your hardware revision. Confirm the waveform with the actual X300 attached in Section 7; if it collapses, stop and resolve the input load or use an X300-compatible PPS distribution driver. Do not raise the supply above 5 V to compensate.')

page('7 Connect the cable to the X300')
step(15,'After the bench checks pass, turn off the F9P, buffer and X300. Keep your PPS cable clearly labelled at both ends.')
step(16,'Find PPS/TRIG IN on the X300 rear panel by reading the printed text. Screw the SMA male cable onto this socket. Start the thread by hand, keep the connector straight and turn the coupling nut rather than twisting the coax.')
step(17,'Leave PPS/TRIG OUT, REF OUT and GPS unused for this PPS connection. If another PPS source is already connected to PPS/TRIG IN, replace that cable; never combine two outputs with a passive T.')
step(18,'If you already use a PRS10 or other suitable 10 MHz source, retain its separate cable into REF IN. PPS and the frequency reference have different jobs. If no external clock is connected, use the X300 internal clock for this PPS input test.')
step(19,'For the loaded check, use a short SMA T at PPS/TRIG IN and a high-impedance probe adapter on the spare branch. Keep the probe branch very short. Do not add a 50 ohm scope termination or leave a long open cable stub.')
step(20,'Power the X300, then the buffer, then the F9P. Wait for valid receiver timing. Measure at the X300 connector with the cable attached: confirm a clean, near-5 V pulse, near-zero low level, and one rising edge each second.')
step(21,'If the loaded waveform is reduced, excessively slow, or rings through the threshold, stop here. Check the ground, connector, load and buffer. Do not treat a blinking LED as a substitute for the loaded waveform check.')
step(22,'When the measurement passes, power down before removing the measurement T and connecting the final direct cable. Restore power in the same order and run Section 8. For shutdown, turn the F9P off before the buffer, then power off the X300 as appropriate.')
h('Final physical connection list')
table(['From','To'],[
('F9P PPS pad','TC4427A pin 2, IN A.'),
('F9P GND pad','Buffer ground and supply negative.'),
('TC4427A pin 7, OUT A','47 ohm series resistor, then coax centre.'),
('Buffer ground','Coax shield and SMA shell.'),
('Far-end SMA male','X300 PPS/TRIG IN.'),
('Separate 10 MHz source, if used','X300 REF IN.'),
('F9P USB-C','Host computer for power, configuration and time messages.')],[2.75,4.02])
p('Avoid hanging the coax from the F9P pads. Mount the buffer close to the receiver, anchor the cable and maintain a short ground return. Keep this timing cable separate from the receiver antenna feed.')

page('8 Check that UHD detects the PPS')
p('Run this section only while the X300 is idle. The UHD test below changes device time. It checks PPS reception and latching; it does not measure electrical quality or assign the real GPS epoch. No RF transmission is needed.')
p('The commands use the previously used lab address 192.168.10.2. Confirm that this is still your X300 address. The local example executable was found at the path below when this guide was prepared; no hardware test was run for this guide.')
h('Test with an existing external 10 MHz source')
code('rtk proxy /opt/homebrew/lib/uhd/examples/test_pps_input \\\n  --args "addr=192.168.10.2,clock_source=external" \\\n  --source external')
p('Use this only when the 10 MHz source is connected and operating. Expect the example to detect successive PPS edges and finish with Success!. Stop if it reports no PPS or a reference-clock problem.')
h('Test without an external 10 MHz source')
code('rtk proxy /opt/homebrew/lib/uhd/examples/test_pps_input \\\n  --args "addr=192.168.10.2,clock_source=internal" \\\n  --source external')
p('Choose one of the two tests. Both use the external PPS input; only the sample-clock source differs. On another computer, use its UHD example path and omit the rtk proxy wrapper if RTK is not installed.')
h('Select the same sources in the real application')
code('usrp.set_clock_source("external")  # if a 10 MHz source is connected\n# Otherwise: usrp.set_clock_source("internal")\nusrp.set_time_source("external")')
p('These lines show the UHD API intent. Source settings must be made in the actual application session; do not assume a preceding test configures every later program. [1]')
h('What success means')
p('You have verified that the chosen cable and input can deliver and detect the pulse. For absolute GPS timestamps, software must still associate a specific pulse with its GPS week and time of week and latch the corresponding device time. USB arrival time and set_time_now(0) do not establish that association. [5]')
p('A GPS-derived PPS also does not automatically discipline an independent 10 MHz source. Keep the roles of pulse alignment, oscillator frequency and epoch labelling separate.')
h('Record the physical delay')
p('Measure the delay from the F9P rising edge to the edge at the X300 connector, including the buffer and cable. Record cable lengths and any correction. Do not subtract a generic gate-delay number and claim nanosecond calibration; supply, load, probe setup and cable affect the result.')

page('9 Troubleshooting and acceptance record')
table(['Observation','Next check'],[
('No PPS at F9P pad','Verify USB power, sky reception, time validity and timepulse read-back. Confirm you probed PPS, not RTK or RST.'),
('PPS exists but buffer stays low','Verify +5 V on pin 6, GND on pin 3, signal on pin 2, pin orientation and input amplitude above the driver threshold.'),
('Pulse is inverted','Check the IC marking. This circuit uses TC4427A, not TC4426A, and OUT A on pin 7.'),
('Good pulse unloaded, poor pulse at X300','Check for an added 50 ohm termination, an unexpected input load, long probe stub, inadequate buffer supply or wrong resistor.'),
('Scope shows PPS but UHD reports none','Check PPS/TRIG IN, common ground, external time source selection, pulse amplitude and that the correct X300 is addressed.'),
('Pulse continues after antenna removal','Check receiver validity and holdover behaviour in messages; do not infer live GPS lock from pulse presence.'),
('Reference-clock error','Check the separate 10 MHz cable and clock source selection. PPS cannot supply the missing frequency reference.'),
('UHD test passes but GPS time is wrong','Check epoch association and time-grid interpretation in software. The test establishes latching only.')],[2.35,4.42],9.4)
h('Fill in after the physical installation')
table(['Record','Measured value or result'],[
('Receiver firmware and carrier identification',''),
('F9P PPS low / high voltage',''),
('Pulse period / high width',''),
('Buffer supply voltage / part marking',''),
('Loaded PPS low / high voltage at X300',''),
('Cable type / length / measured added delay',''),
('GNSS time-valid state and time grid',''),
('X300 address / hardware revision',''),
('UHD PPS test result / date','')],[3.60,3.17],9.4)
p('Ready for integration when: solder joints and strain relief are sound; PPS is valid at the receiver; the loaded input waveform passes; and the idle UHD PPS test succeeds. Record absolute GPS epoch binding separately.')

page('10 Sources and design limits')
p('The photographs establish visible board labels. Manufacturer documents establish receiver and interface capabilities. The short-wire lengths, buffer assembly, resistor value and measurement sequence are engineering proposals for this bench setup, to be qualified on the actual hardware.')
link('[1] Ettus UHD manual for USRP X3x0 series','https://files.ettus.com/manual/page_usrp_x3x0.html')
p('Rear connector names, 5 Vpp PPS guidance, separate reference-clock input, and UHD clock/time source selection. See the Rear Panel and PPS sections.')
link('[2] u-blox ZED-F9P-04B data sheet UBX-21044850','https://content.u-blox.com/sites/default/files/ZED-F9P-04B_DataSheet_UBX-21044850.pdf')
p('TIMEPULSE pin assignment and electrical limits, including the 4 mA TIMEPULSE drive/sink specification. The module supply and GPIO specifications must not be confused with the carrier USB supply.')
link('[3] u-blox F9 HPG 1.32 interface description UBX-22008968','https://content.u-blox.com/sites/default/files/documents/u-blox-F9-HPG-1.32_InterfaceDescription_UBX-22008968.pdf')
p('CFG-TP configuration, pages 268-270. The example settings are for the documented interface; read the receiver firmware version before applying them.')
link('[4] Microchip TC4426A TC4427A TC4428A data sheet DS20001423K','https://ww1.microchip.com/downloads/aemDocuments/documents/APID/ProductDocuments/DataSheets/TC4426A-TC4427A-TC4428A-1.5A-Dual-High-Speed-Power-MOSFET-Drivers-20001423.pdf')
p('TC4427A non-inverting function; 8-pin package pin map; 2.4 V minimum high and 0.8 V maximum low input levels; supply and input-voltage limits. Only the TC4427A non-inverting function is used in this guide.')
link('[5] u-blox ZED-F9P integration manual UBX-18010802','https://content.u-blox.com/sites/default/files/ZED-F9P_IntegrationManual_UBX-18010802.pdf')
p('Timepulse configuration and timing-message interpretation. UBX-TIM-TP relates to the next pulse. Timing validity, time-grid selection and delay calibration remain part of integration.')
h('What has and has not been verified')
p('The PPS/GND labels were checked in the supplied F9P photo. Manufacturer documents were reviewed, and the generated annotations were inspected. The document was rendered and checked for layout. No soldering, voltage measurement, X300 load measurement, GNSS status read-back or physical PPS test was performed for this guide.')
p('The illustrated USRP panel is generated, with deliberately schematic port positions. It is not an edited photograph of your X300. Match PPS/TRIG IN on the physical unit, regardless of where a connector appears in the illustration.')
p('The F9P carrier schematic and exact X300 input load were not available for this hardware. Accordingly, the buffered design includes a mandatory loaded-waveform check. If you prefer a direct PPS-to-SMA connection, first establish the carrier output voltage/drive, X300 input threshold/load and waveform quality; the existence of a PPS pad alone does not settle compatibility.')
h('Image files')
p('images/f9p-annotated.png - annotation of your supplied F9P photograph.\nimages/x300-pps-connectors.png - generated connector identification illustration.\nimage-prompts.md - the exact prompts used with the built-in GPT Image tool.')

for el in doc.styles.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
for el in doc.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
doc.save(OUT/'zed-f9p-to-x300-pps-wiring.docx')
print(OUT/'zed-f9p-to-x300-pps-wiring.docx')
